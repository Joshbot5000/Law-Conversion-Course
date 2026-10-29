"""Builds the made-up test data used by the browser tests.

Nothing here comes from a real library, textbook or lecture notes: the case names are
public, but every summary, quote and source passage is short sample text written for testing.

Run:  python3 tests/fixtures/make_fixtures.py
Writes: tests/sample_seed.js, tests/fixtures/sample-import.json, tests/fixtures/sample-rules.json,
        tests/fixtures/sample-chapter.pdf, tests/fixtures/sample-notes.docx
"""
import json, os

HERE = os.path.dirname(os.path.abspath(__file__))
TESTS = os.path.dirname(HERE)
STAMP = '2026-01-10T09:00:00Z'


def hash32(s):
    # Same FNV-1a hash the app uses to choose a storage shard (js_core.js hash32), over UTF-16 code units.
    b = s.encode('utf-16-le')
    h = 0x811c9dc5
    for k in range(0, len(b), 2):
        h ^= b[k] | (b[k + 1] << 8)
        h = (h * 0x01000193) & 0xFFFFFFFF
    return h


def slug(s):
    out = ''.join(c.lower() if c.isascii() and c.isalnum() else '-' for c in s)
    while '--' in out:
        out = out.replace('--', '-')
    return out.strip('-')[:60]


# ---- Sources (made-up sample passages; form feeds mark page breaks) ----
SOURCES = {
    'src-sample-contract': ('Sample contract chapter: Agreement and Estoppel', 'contract', [
        'Agreement\n\nAn offer is a statement of the terms on which the offeror is willing to be bound. '
        'In Carlill v Carbolic Smoke Ball Co the advert was treated as an offer to the whole world because it showed a clear intention to be bound. '
        'By contrast, in Partridge v Crittenden an advert for birds was only an invitation to treat. '
        'Goods on display in a shop window are an invitation to treat, as Fisher v Bell shows.',
        'Acceptance\n\nAcceptance must mirror the offer. A counter-offer destroys the original offer, as in Hyde v Wrench. '
        'Silence cannot normally amount to acceptance, which is the point of Felthouse v Bindley. '
        'Under the postal rule an acceptance takes effect when posted, following Adams v Lindsell.',
        'Promissory estoppel\n\nWhere a promise not to enforce strict legal rights is relied on, the promisor may be prevented from going back on it. '
        'The doctrine grew from Hughes v Metropolitan Railway and was restated in Central London Property Trust v High Trees House. '
        'Combe v Combe confirms that promissory estoppel is a shield and not a sword. '
        'Terms may be implied from custom, as in Hutton v Warren, where a local custom gave the tenant an allowance.',
    ]),
    'src-sample-tort': ('Sample tort chapter: Negligence and Causation', 'tort', [
        'Duty of care\n\nDonoghue v Stevenson established that a manufacturer owes a duty to the ultimate consumer. '
        'In novel situations the courts reason by analogy with decided cases, as explained in Robinson v Chief Constable of West Yorkshire.',
        'Causation\n\nThe claimant must show that the breach caused the damage. In Barnett v Chelsea and Kensington Hospital the claimant failed the but-for test. '
        'In Bonnington Castings v Wardlaw the dust from the guilty source made a material contribution to the pneumoconiosis. '
        'Where the science cannot separate the causes, a material contribution to the harm is enough.',
    ]),
    'src-sample-criminal': ('Sample criminal chapter: Mens Rea', 'criminal', [
        'Intention\n\nOblique intention is found where death or serious injury was a virtual certainty and the defendant appreciated that, following R v Woollin.',
        'Recklessness\n\nSubjective recklessness requires awareness of a risk which it was unreasonable to take, as stated in R v G.',
    ]),
}

# ---- Cards: (module, topic, type, name, citation, facts, ratio, strategy, sourceId, quote, page) ----
C = []
def card(module, topic, name, ratio, facts='', strategy='', typ='case', citation='', src='', quote='', page=0, cid=None):
    C.append(dict(module=module, topic=topic, type=typ, name=name, citation=citation, facts=facts, ratio=ratio,
                  strategy=strategy, sourceId=src, sourceQuote=quote, sourceRef=('p. %d' % page if page else ''), cid=cid))

sc = 'src-sample-contract'
card('contract', 'Formation – Offer', 'Carlill v Carbolic Smoke Ball Co', 'An advert can be a unilateral offer when it shows a clear intention to be bound, such as depositing money to show sincerity.',
     'A company promised a reward to anyone who caught flu after using its product as directed.', 'Use for reward adverts and unilateral offers.', citation='[1893] 1 QB 256 (CA)',
     src=sc, quote='the advert was treated as an offer to the whole world because it showed a clear intention to be bound', page=1)
card('contract', 'Formation – Offer', 'Partridge v Crittenden', 'An advert for goods is usually an invitation to treat rather than an offer.',
     'A seller advertised birds for sale in a magazine.', 'Contrast with Carlill when an advert is in issue.', citation='[1968] 1 WLR 1204',
     src=sc, quote='an advert for birds was only an invitation to treat', page=1)
card('contract', 'Formation – Offer', 'Fisher v Bell', 'Goods displayed in a shop window are an invitation to treat.',
     'A flick knife was displayed in a shop window with a price ticket.', 'Shop displays: no offer until the customer offers to buy.', citation='[1961] 1 QB 394',
     src=sc, quote='Goods on display in a shop window are an invitation to treat', page=1)
card('contract', 'Formation – Offer', 'Fisher v Bell', 'A shop window display is an invitation to treat, not an offer.',
     'Knife displayed in a window.', 'Duplicate copy kept for testing the duplicate finder.', citation='[1961] 1 QB 394')
card('contract', 'Formation – Offer', 'Pharmaceutical Society v Boots Cash Chemists', 'Goods on self-service shelves are an invitation to treat; the customer makes the offer at the till.',
     'Medicines were sold from open shelves under a pharmacist’s supervision at the till.', 'Pair with Fisher v Bell.', citation='[1953] 1 QB 401 (CA)')
card('contract', 'Formation – Acceptance', 'Hyde v Wrench', 'A counter-offer rejects and destroys the original offer.',
     'A buyer offered a lower price, then tried to accept the original offer.', 'Check every reply: acceptance or counter-offer?', citation='(1840) 3 Beav 334',
     src=sc, quote='A counter-offer destroys the original offer, as in Hyde v Wrench', page=2)
card('contract', 'Formation – Acceptance', 'Felthouse v Bindley', 'Silence cannot normally amount to acceptance.',
     'An uncle said he would treat a horse as his unless he heard otherwise.', 'Use when an offeror tries to impose acceptance by silence.', citation='(1862) 11 CB (NS) 869',
     src=sc, quote='Silence cannot normally amount to acceptance', page=2)
card('contract', 'Formation – Acceptance', 'Adams v Lindsell', 'Under the postal rule, acceptance takes effect on posting.',
     'A letter of acceptance was delayed in the post.', 'Apply only where post is a reasonable method of acceptance.', citation='(1818) 1 B & Ald 681')
card('contract', 'Formation – Acceptance', 'Entores v Miles Far East Corp', 'Acceptance by instantaneous communication takes effect when received.',
     'Acceptance was sent by telex between countries.', 'Contrast instantaneous methods with the postal rule.', citation='[1955] 2 QB 327 (CA)')
card('contract', 'Promissory Estoppel', 'Hughes v Metropolitan Railway', 'Where a party leads another to believe strict rights will not be enforced, equity may prevent their sudden enforcement.',
     'A landlord opened negotiations and then demanded repairs within the original deadline.', 'The historical root of promissory estoppel.', citation='(1877) 2 App Cas 439 (HL)')
card('contract', 'Promissory Estoppel', 'Central London Property Trust v High Trees House', 'A promise to accept less, relied on, may be binding in equity while the circumstances last.',
     'A landlord halved the rent during wartime when flats were empty.', 'Use for promises to accept part payment of a debt.', citation='[1947] KB 130',
     src=sc, quote='restated in Central London Property Trust v High Trees House', page=3)
card('contract', 'Promissory Estoppel', 'Combe v Combe', 'Promissory estoppel is a shield and not a sword; it cannot create a new cause of action.',
     'A wife sued on her former husband’s promise of maintenance.', 'Use to stop a claimant founding a claim on estoppel alone.', citation='[1951] 2 KB 215 (CA)',
     src=sc, quote='Combe v Combe confirms that promissory estoppel is a shield and not a sword', page=3)
card('contract', 'Promissory Estoppel', 'Combe v Combe', 'Estoppel can only be used defensively.',
     'Maintenance promise after divorce.', 'Second copy kept for testing duplicate merging.', citation='[1951] 2 KB 215')
card('contract', 'Promissory Estoppel', 'Promissory estoppel — the four conditions',
     'There must be a clear promise not to enforce strict rights, reliance on it, inequity in going back on it, and it acts as a shield, not a sword.',
     'A landlord promises to accept lower rent and the tenant relies on that promise.', 'Work through each condition in turn.', typ='principle',
     citation='Key authorities: Hughes v Metropolitan Railway; Central London Property Trust v High Trees House; Combe v Combe')
card('contract', 'Consideration', 'Currie v Misa', 'Consideration is some right, interest, profit or benefit to one party, or detriment to the other.',
     'A dispute about whether a bill of exchange was given for value.', 'Use for the classic definition.', citation='(1875) LR 10 Ex 153')
card('contract', 'Consideration', 'Chappell & Co v Nestlé', 'Consideration must be sufficient but need not be adequate; even wrappers can count.',
     'Chocolate wrappers were required to buy a record at a reduced price.', 'Use where the consideration looks trivial.', citation='[1960] AC 87 (HL)')
card('contract', 'Consideration', 'Stilk v Myrick', 'Performing an existing contractual duty owed to the promisor is not good consideration.',
     'Sailors were promised extra wages for sailing a ship home short-handed.', 'Starting point for promises of extra pay.', citation='(1809) 2 Camp 317')
card('contract', 'Consideration', 'Williams v Roffey Bros', 'A practical benefit to the promisor can be good consideration for a promise of extra pay.',
     'A main contractor promised a carpenter more money to finish on time.', 'Distinguish Stilk v Myrick on practical benefit.', citation='[1991] 1 QB 1 (CA)')
card('contract', 'Terms – Implied by the Courts', 'Hutton v Warren', 'A term may be implied from local custom where the contract is silent.',
     'An outgoing tenant claimed a customary allowance for seed and labour.', 'Use for terms implied by custom.', citation='(1836) 1 M & W 466',
     src=sc, quote='Terms may be implied from custom, as in Hutton v Warren', page=3, cid='contract-ch3-hutton-v-warren')
card('contract', 'Terms – Implied by the Courts', 'The Moorcock', 'A term may be implied in fact where it is needed to give the contract business efficacy.',
     'A ship was damaged when grounding at a jetty at low tide.', 'Business efficacy test for implied terms.', citation='(1889) 14 PD 64 (CA)')
card('contract', 'Terms – Implied by the Courts', 'Liverpool City Council v Irwin', 'Terms may be implied in law into a type of contract where necessary.',
     'Tenants complained about broken lifts and stairwell lighting.', 'Use for terms implied in law.', citation='[1977] AC 239 (HL)')
card('contract', 'Misrepresentation', 'Howard Marine & Dredging Co Ltd v A Ogden & Sons (Excavations) Ltd', 'Under the Misrepresentation Act 1967, s 2(1), the representor must prove reasonable grounds for believing the statement.',
     'Barge capacity was wrongly stated during negotiations.', 'Use for the burden of proof under s 2(1).', citation='[1978] QB 574 (CA)')
card('contract', 'Misrepresentation', 'Bisset v Wilkinson', 'A genuinely held opinion is not a statement of fact.',
     'A seller gave an estimate of how many sheep land could carry.', 'Use to separate opinion from fact.', citation='[1927] AC 177 (PC)')
card('contract', 'Misrepresentation', 'Edgington v Fitzmaurice', 'A statement of present intention can be a misrepresentation if the intention was not held.',
     'Company directors misstated why they were raising money.', 'Use for statements of intention.', citation='(1885) 29 Ch D 459 (CA)')
card('contract', 'Frustration', 'Taylor v Caldwell', 'A contract is frustrated where the subject matter is destroyed without fault.',
     'A music hall burned down before the concerts were due.', 'Classic destruction example.', citation='(1863) 3 B & S 826')
card('contract', 'Frustration', 'Krell v Henry', 'A contract can be frustrated where its common purpose fails.',
     'Rooms were hired to watch a coronation procession that was cancelled.', 'Contrast with Herne Bay.', citation='[1903] 2 KB 740 (CA)')
card('contract', 'Frustration', 'Davis Contractors v Fareham UDC', 'Mere hardship or expense does not frustrate a contract.',
     'Builders faced labour shortages and higher costs.', 'Use to reject frustration claims based on cost.', citation='[1956] AC 696 (HL)')
card('contract', 'Exclusion Clauses', "L'Estrange v Graucob", 'A party who signs a contractual document is bound by it, whether or not they read it.',
     'A café owner signed an order form containing an exclusion clause.', 'Incorporation by signature.', citation='[1934] 2 KB 394')
card('contract', 'Exclusion Clauses', 'Olley v Marlborough Court Hotel', 'Terms introduced after the contract is made are not incorporated.',
     'A notice in a hotel bedroom tried to exclude liability.', 'Check timing of notice.', citation='[1949] 1 KB 532 (CA)')
card('contract', 'Exclusion Clauses', 'Thornton v Shoe Lane Parking', 'Onerous terms need to be fairly brought to the other party’s attention.',
     'A ticket from a car park machine referred to conditions inside.', 'Red hand rule for unusual terms.', citation='[1971] 2 QB 163 (CA)')
card('contract', 'Intention to Create Legal Relations', 'Balfour v Balfour', 'Domestic agreements are presumed not to be legally binding.',
     'A husband promised his wife a monthly allowance while abroad.', 'Domestic presumption.', citation='[1919] 2 KB 571 (CA)')
card('contract', 'Intention to Create Legal Relations', 'Merritt v Merritt', 'The domestic presumption is rebutted where spouses have separated.',
     'A separated husband promised to transfer the house to his wife.', 'Rebutting the presumption.', citation='[1970] 1 WLR 1211 (CA)')

st = 'src-sample-tort'
card('tort', 'Negligence – Duty of Care', 'Donoghue v Stevenson',
     'A manufacturer owes a duty of care to the ultimate consumer of its products, and you must take reasonable care to avoid acts or omissions you can reasonably foresee would injure your neighbour.',
     'A woman fell ill after drinking ginger beer from an opaque bottle that was said to contain a decomposed snail.',
     'Start any duty question here, then ask whether the situation is covered by an established category.',
     citation='[1932] AC 562 (HL)', src=st, quote='Donoghue v Stevenson established that a manufacturer owes a duty to the ultimate consumer', page=1)
card('tort', 'Negligence – Duty of Care', 'Caparo Industries v Dickman', 'In novel cases, consider foreseeability, proximity and whether a duty would be fair, just and reasonable.',
     'Investors relied on audited accounts when buying shares.', 'Use incrementally for genuinely novel situations.', citation='[1990] 2 AC 605 (HL)')
card('tort', 'Negligence – Duty of Care', 'Robinson v Chief Constable of West Yorkshire', 'Courts apply established principles and develop the law incrementally by analogy.',
     'A passer-by was injured during a street arrest.', 'Use to explain when Caparo is not needed.', citation='[2018] UKSC 4',
     src=st, quote='the courts reason by analogy with decided cases', page=1)
card('tort', 'Negligence – Breach of Duty', 'Bolton v Stone', 'A very small risk may be one a reasonable person would not guard against.',
     'A cricket ball was hit out of the ground, which had happened rarely.', 'Weigh probability of harm.', citation='[1951] AC 850 (HL)')
card('tort', 'Negligence – Breach of Duty', 'Paris v Stepney Borough Council', 'Where the harm would be especially serious, more precautions are expected.',
     'A worker already blind in one eye was not given goggles.', 'Seriousness of harm.', citation='[1951] AC 367 (HL)')
card('tort', 'Negligence – Causation', 'Barnett v Chelsea and Kensington Hospital', 'The breach must be a but-for cause of the damage.',
     'A man sent home from hospital died of poisoning that could not have been treated in time.', 'Apply the but-for test first.', citation='[1969] 1 QB 428',
     src=st, quote='the claimant failed the but-for test', page=2)
card('tort', 'Negligence – Causation', 'Bonnington Castings v Wardlaw', 'A material contribution to the harm is enough where the causes cannot be separated.',
     'A worker developed lung disease from dust from two sources, only one in breach.', 'Use for cumulative causes.', citation='[1956] AC 613 (HL)',
     src=st, quote='In Bonnington Castings v Wardlaw the dust from the guilty source made a material contribution to the pneumoconiosis', page=2)
card('tort', 'Negligence – Causation', 'McGhee v National Coal Board', 'Materially increasing the risk of harm can be treated as causing it in limited cases.',
     'A worker developed dermatitis after no washing facilities were provided.', 'Contrast with Bonnington.', citation='[1973] 1 WLR 1 (HL)')
card('tort', 'Employers’ Liability', 'Wilsons & Clyde Coal Co v English', 'An employer owes a personal, non-delegable duty to take reasonable care for employees’ safety.',
     'A miner was crushed by machinery while leaving the pit.', 'Safe system, plant, place and competent staff.', citation='[1938] AC 57 (HL)')
card('tort', 'Employers’ Liability', 'Latimer v AEC', 'An employer need only take reasonable precautions, not eliminate every risk.',
     'A factory floor was slippery after flooding; sawdust was spread but not everywhere.', 'Balance risk against cost.', citation='[1953] AC 643 (HL)')

sk = 'src-sample-criminal'
card('criminal', 'Mens Rea – Intention', 'R v Woollin', 'Intention may be found where death or serious injury was a virtual certainty and the defendant appreciated this.',
     'A father threw his baby onto a hard surface.', 'Use for oblique intention.', citation='[1999] 1 AC 82 (HL)',
     src=sk, quote='death or serious injury was a virtual certainty and the defendant appreciated that', page=1)
card('criminal', 'Mens Rea – Intention', 'R v Woollin', 'Virtual certainty test for oblique intention.', 'Baby thrown in anger.', 'Duplicate copy for testing.', citation='[1999] 1 AC 82')
card('criminal', 'Mens Rea – Recklessness', 'R v G', 'Recklessness is subjective: the defendant must be aware of a risk that it was unreasonable to take.',
     'Two boys set fire to newspapers behind a shop and the fire spread.', 'Use the subjective test.', citation='[2003] UKHL 50',
     src=sk, quote='Subjective recklessness requires awareness of a risk which it was unreasonable to take', page=2)
card('criminal', 'Mens Rea – Recklessness', 'R v Cunningham', 'Malice means intention or subjective recklessness.',
     'A man tore a gas meter from a wall and gas seeped next door.', 'Origin of subjective recklessness.', citation='[1957] 2 QB 396 (CA)')
card('criminal', 'Actus Reus – Omissions', 'R v Miller', 'A person who creates a dangerous situation is under a duty to act once aware of it.',
     'A squatter fell asleep with a cigarette and did nothing when he woke to a fire.', 'Duty arising from creating danger.', citation='[1983] 2 AC 161 (HL)')
card('criminal', 'Causation', 'R v White', 'The conduct must be a factual (but-for) cause of the result.',
     'A man poisoned his mother, who died of a heart attack before the poison worked.', 'Factual causation first.', citation='[1910] 2 KB 124 (CA)')
card('criminal', 'Causation', 'R v Pagett', 'A reasonable act of self-defence by a third party does not break the chain of causation.',
     'A man used a girl as a shield when police fired back.', 'Intervening acts.', citation='(1983) 76 Cr App R 279 (CA)')

card('els', 'Statutory Interpretation', 'Pepper v Hart', 'Hansard may be used where legislation is ambiguous and a clear ministerial statement resolves it.',
     'Teachers disputed tax on a concessionary school fees benefit.', 'Use for external aids.', citation='[1993] AC 593 (HL)')
card('els', 'Statutory Interpretation', 'Smith v Hughes', 'The mischief rule looks at the problem the statute was meant to remedy.',
     'Women solicited from balconies and windows rather than in the street.', 'Use with the mischief rule.', citation='[1960] 1 WLR 830')
card('public', 'Parliamentary Sovereignty', 'R (Jackson) v Attorney General', 'The Parliament Acts procedure produces valid primary legislation.',
     'The Hunting Act 2004 was passed using the Parliament Acts.', 'Use for manner and form arguments.', citation='[2005] UKHL 56')
card('public', 'Royal Prerogative', 'R (Miller) v Secretary of State for Exiting the EU', 'The prerogative cannot be used to change domestic law or remove rights.',
     'The government planned to give notice to leave the EU without an Act.', 'Limits of the prerogative.', citation='[2017] UKSC 5')

card('public', 'Parliamentary Sovereignty', 'Thoburn v Sunderland City Council', 'Constitutional statutes are not impliedly repealed by later Acts.',
     'Traders were prosecuted for selling goods in imperial measures only.', 'Use for implied repeal and constitutional statutes.', citation='[2002] EWHC 195 (Admin)')
card('public', 'Royal Prerogative', 'Case of Proclamations', 'The King cannot create new offences by proclamation; the prerogative is limited by law.',
     'Proclamations tried to stop new building in London and the making of starch.', 'Use for the legal limits of the prerogative.', citation='(1611) 12 Co Rep 74')

# Pad with clearly invented cards so the library is big enough for review and paging tests.
TOPICS = {'contract': 'Sample Topic – Remedies', 'tort': 'Sample Topic – Nuisance', 'criminal': 'Sample Topic – Theft', 'land': 'Sample Topic – Easements'}
for i in range(1, 25):
    m = list(TOPICS)[i % 4]
    card(m, TOPICS[m], 'Example Claimant %d v Sample Defendant' % i, 'Sample rule number %d, written only for testing the app.' % i,
         'Invented facts for test card %d.' % i, 'Invented strategy %d.' % i)


def build():
    seed, entries, used = {}, [], set()
    for i, c in enumerate(C):
        cid = c.pop('cid') or (c['module'] + '-' + slug(c['name']))
        while cid in used:
            cid += '-2'
        used.add(cid)
        e = dict(id=cid, tags=[], memorised=False, verified=False, createdAt=STAMP, updatedAt=STAMP, **c)
        entries.append(e)
        seed.setdefault('lib/s%02d' % (hash32(cid) % 32), {'e': {}})['e'][cid] = e
    for sid, (title, module, pages) in SOURCES.items():
        text = '\f'.join(pages)
        seed['src/' + sid] = dict(id=sid, title=title, module=module, kind='pdf', pages=len(pages), chars=len(text), parts=1, createdAt=STAMP)
        seed['srctext/' + sid + '_0'] = dict(sid=sid, n=0, text=text)
    js = ('// Made-up sample library for the tests (generated by fixtures/make_fixtures.py). Not real study notes.\n'
          'window.__SEED = %s;\nwindow.__SEED_COUNT = %d;\nwindow.__SEED_SOURCES = %d;\n') % (json.dumps(seed, ensure_ascii=False, indent=0), len(entries), len(SOURCES))
    open(os.path.join(TESTS, 'sample_seed.js'), 'w').write(js)

    # v1 backup files for the offline import test (new ids, so they add to the 3 built-in samples)
    def backup(items):
        return dict(app='Law Revision Repository', format='law-revision-backup', version=1, exportedAt=STAMP, count=len(items), entries=items)
    imp = [dict(id='import-%d' % i, module='contract', topic='Imported – Sample', type='case', name='Imported Claimant %d v Defendant' % i,
                citation='', facts='Invented facts %d.' % i, ratio='Invented rule %d for import testing.' % i, strategy='', tags=['sample'],
                memorised=False, createdAt=STAMP, updatedAt=STAMP) for i in range(1, 21)]
    rules = [dict(id='rule-%d' % i, module='tort', topic='Imported – Rules', type='principle', name='Sample principle %d' % i,
                  citation='', facts='', ratio='Invented principle %d for import testing.' % i, strategy='', tags=['sample'],
                  memorised=False, createdAt=STAMP, updatedAt=STAMP) for i in range(1, 11)]
    json.dump(backup(imp), open(os.path.join(HERE, 'sample-import.json'), 'w'), indent=1)
    json.dump(backup(rules), open(os.path.join(HERE, 'sample-rules.json'), 'w'), indent=1)

    # A short made-up "chapter" PDF for the add-notes test
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    pdf = canvas.Canvas(os.path.join(HERE, 'sample-chapter.pdf'), pagesize=A4)
    pages = [
        ['Sample Chapter: Frustration', '', 'A contract may be frustrated when performance becomes impossible.',
         'In Taylor v Caldwell the hall burned down before the concerts could take place.',
         'In Krell v Henry the purpose of hiring the rooms failed when the procession was cancelled.'],
        ['Limits of frustration', '', 'In Herne Bay Steam Boat Co v Hutton the cruise still had some purpose, so there was no frustration.',
         'In Davis Contractors v Fareham UDC extra cost and delay did not frustrate the building contract.',
         'Self-induced frustration does not count, as Maritime National Fish v Ocean Trawlers shows.'],
        ['Consequences', '', 'The Law Reform (Frustrated Contracts) Act 1943 deals with money paid before the event.',
         'In Gamerco v ICM the court used its discretion under the Act after a stadium was found unsafe.'],
    ]
    for lines in pages:
        y = 780
        for ln in lines:
            pdf.drawString(60, y, ln); y -= 22
        pdf.showPage()
    pdf.save()

    # A short made-up Word document for the .docx reading test
    import docx
    d = docx.Document()
    d.add_heading('Sample tort notes', 1)
    for i in range(1, 31):
        d.add_heading('Sample topic %d' % i, 2)
        d.add_paragraph('Invented note %d: the duty of care is assessed by analogy with decided cases such as Donoghue v Stevenson. ' % i * 6)
    d.save(os.path.join(HERE, 'sample-notes.docx'))
    print('cards:', len(entries), 'sources:', len(SOURCES))


if __name__ == '__main__':
    build()
